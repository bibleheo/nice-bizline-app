"""결과 엑셀 출력 - 결과/미발견·오류/확인필요/실행로그 4개 시트."""
from __future__ import annotations

from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


_BASE_HEADERS_BEFORE_FINANCE = [
    "회사명", "대표자", "사업자번호", "주소", "업종", "설립일", "대표번호", "종업원수",
    "휴폐업정보",
]
_BASE_HEADERS_AFTER_FINANCE = ["신용등급", "조회상태", "조회일시", "비고"]


def build_headers(records: list[dict], finance_years: int = 1) -> list[str]:
    """records와 옵션을 보고 헤더 리스트를 구성.

    단년: 매출액(백만원)/영업이익(백만원)/당기순이익(백만원)
    다년: records에 등장하는 연도별 컬럼을 수집해 최신연도부터 정렬
    """
    if finance_years <= 1:
        finance = ["매출액(백만원)", "영업이익(백만원)", "당기순이익(백만원)"]
    else:
        # 모든 record에서 "매출액(YYYY)" / "영업이익(YYYY)" / "당기순이익(YYYY)" 키 수집
        years: set[str] = set()
        for rec in records:
            for k in rec:
                for prefix in ("매출액(", "영업이익(", "당기순이익("):
                    if k.startswith(prefix) and k.endswith(")"):
                        years.add(k[len(prefix):-1])
        sorted_years = sorted(years, reverse=True)
        finance = []
        for metric in ("매출액", "영업이익", "당기순이익"):
            for y in sorted_years:
                finance.append(f"{metric}({y})")
    # 결산일자: 재무 값이 어느 결산 기준인지 표시
    return (_BASE_HEADERS_BEFORE_FINANCE + ["결산일자"] + finance
            + _BASE_HEADERS_AFTER_FINANCE)


# 단년 모드 호환용 (테스트/외부 참조)
HEADERS = build_headers([], finance_years=1)

_STATUS_COLORS = {
    "성공": "D9EAD3",
    "미발견": "FCE5CD",
    "확인필요": "FFF2CC",
    "오류": "F4CCCC",
}


def _thin():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def _header(cell, text):
    cell.value = text
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", start_color="2F5496")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _thin()


_COL_WIDTHS = {
    "회사명": 22, "대표자": 12, "사업자번호": 16, "주소": 32, "업종": 22,
    "설립일": 12, "대표번호": 14, "종업원수": 10, "휴폐업정보": 12, "결산일자": 12,
    "신용등급": 10,
    "조회상태": 10, "조회일시": 18, "비고": 30,
}


def _col_width(h: str) -> int:
    if h in _COL_WIDTHS:
        return _COL_WIDTHS[h]
    if h.startswith(("매출액", "영업이익", "당기순이익")):
        return 14
    return 14


def _is_numeric_column(h: str) -> bool:
    if h == "종업원수":
        return True
    return h.startswith(("매출액", "영업이익", "당기순이익"))


def _data(cell, row):
    cell.font = Font(size=10)
    cell.fill = PatternFill("solid", start_color="F5F5F5" if row % 2 == 0 else "FFFFFF")
    cell.alignment = Alignment(vertical="center")
    cell.border = _thin()


def _input_columns(records: list[dict]) -> list[str]:
    """records에 연결된 원본 입력 행(_input)의 컬럼을 등장 순서대로 수집."""
    cols: list[str] = []
    for rec in records:
        for k in (rec.get("_input") or {}):
            if k not in cols:
                cols.append(k)
    return cols


def write_results(path: str, records: list[dict], unfound: list[dict],
                  ambiguous: list[dict], summary: dict,
                  finance_years: int = 1) -> None:
    wb = openpyxl.Workbook()
    # 좌측 = 원본 입력 열([입력] 접두), 우측 = 수집 결과 열
    in_cols = _input_columns(records)
    in_headers = [f"[입력] {c}" for c in in_cols]
    headers = build_headers(records, finance_years)
    all_headers = in_headers + headers

    # 시트1: 결과
    ws = wb.active
    ws.title = "결과"
    for col, h in enumerate(all_headers, 1):
        _header(ws.cell(1, col), h)
    ws.row_dimensions[1].height = 22
    for i, h in enumerate(all_headers, 1):
        base = h[5:] if h.startswith("[입력] ") else h
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = _col_width(base)

    for r_idx, rec in enumerate(records, 2):
        # 입력 열: 동명 여러 건이면 첫 행에만 값 표시
        show_input = rec.get("_input_show", True)
        inp = rec.get("_input") or {}
        for c_idx, c in enumerate(in_cols, 1):
            val = inp.get(c, "") if show_input else ""
            cell = ws.cell(r_idx, c_idx, val if val is not None else "")
            _data(cell, r_idx)
        # 결과 열
        off = len(in_cols)
        for c_idx, h in enumerate(headers, 1):
            val = rec.get(h, "")
            cell = ws.cell(r_idx, off + c_idx, val if val is not None else "")
            _data(cell, r_idx)
            if _is_numeric_column(h) and isinstance(val, (int, float)):
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            if h == "조회상태":
                fill = _STATUS_COLORS.get(str(val), "FFFFFF")
                cell.fill = PatternFill("solid", start_color=fill)
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # 시트2: 미발견·오류
    ws2 = wb.create_sheet("미발견·오류")
    for col, h in enumerate(["회사명", "조회상태", "사유"], 1):
        _header(ws2.cell(1, col), h)
    for i, w in enumerate([22, 12, 50], 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for r_idx, item in enumerate(unfound, 2):
        ws2.cell(r_idx, 1, item.get("회사명", ""))
        ws2.cell(r_idx, 2, item.get("조회상태", ""))
        ws2.cell(r_idx, 3, item.get("사유", ""))

    # 시트3: 확인필요
    ws3 = wb.create_sheet("확인필요")
    for col, h in enumerate(["회사명", "채택후보", "다른후보들", "사유"], 1):
        _header(ws3.cell(1, col), h)
    for i, w in enumerate([22, 30, 60, 30], 1):
        ws3.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for r_idx, item in enumerate(ambiguous, 2):
        ws3.cell(r_idx, 1, item.get("회사명", ""))
        ws3.cell(r_idx, 2, item.get("채택후보", ""))
        ws3.cell(r_idx, 3, item.get("다른후보들", ""))
        ws3.cell(r_idx, 4, item.get("사유", ""))

    # 시트4: 실행로그
    ws4 = wb.create_sheet("실행로그")
    for col, h in enumerate(["항목", "값"], 1):
        _header(ws4.cell(1, col), h)
    ws4.column_dimensions["A"].width = 20
    ws4.column_dimensions["B"].width = 30
    rows = [
        ("시작 시각", summary.get("started_at", "")),
        ("종료 시각", summary.get("ended_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))),
        ("총 건수", summary.get("total", 0)),
        ("성공", summary.get("success", 0)),
        ("미발견", summary.get("not_found", 0)),
        ("확인필요", summary.get("ambiguous", 0)),
        ("오류", summary.get("error", 0)),
    ]
    for r_idx, (k, v) in enumerate(rows, 2):
        ws4.cell(r_idx, 1, k)
        ws4.cell(r_idx, 2, v)

    wb.save(path)
