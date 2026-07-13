import openpyxl

from nice_bizline.app.excelio.reader import read_company_list
from nice_bizline.app.excelio.writer import write_results


def _make_input(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)


class TestReader:
    def test_standard_headers(self, tmp_path):
        p = tmp_path / "in.xlsx"
        _make_input(p, [
            ["회사명", "사업자번호", "대표자명"],
            ["삼성전자", "1248100998", "한종희"],
            ["현대자동차", None, None],
        ])
        rows = read_company_list(str(p))
        assert len(rows) == 2
        assert rows[0] == {"회사명": "삼성전자", "사업자번호": "1248100998", "대표자명": "한종희"}
        assert rows[1] == {"회사명": "현대자동차"}

    def test_alias_headers(self, tmp_path):
        """한글 alias (상호/사업자등록번호/대표) 도 인식."""
        p = tmp_path / "in.xlsx"
        _make_input(p, [
            ["상호", "사업자등록번호", "대표"],
            ["X사", "1234567890", "홍길동"],
        ])
        rows = read_company_list(str(p))
        assert rows[0]["회사명"] == "X사"
        assert rows[0]["사업자번호"] == "1234567890"
        assert rows[0]["대표자명"] == "홍길동"

    def test_upche_myeong_alias(self, tmp_path):
        """'업체명' 도 회사명으로 인식."""
        p = tmp_path / "in.xlsx"
        _make_input(p, [
            ["업체명", "사업자번호"],
            ["가나건설", "1112233445"],
        ])
        rows = read_company_list(str(p))
        assert rows[0]["회사명"] == "가나건설"
        assert rows[0]["사업자번호"] == "1112233445"

    def test_address_column_captured(self, tmp_path):
        """'주소'(및 alias '소재지') 컬럼을 캡처해 매칭에 활용."""
        p = tmp_path / "in.xlsx"
        _make_input(p, [
            ["회사명", "소재지"],
            ["한빛엔지니어링", "서울특별시 강남구 테헤란로 1"],
        ])
        rows = read_company_list(str(p))
        assert rows[0]["회사명"] == "한빛엔지니어링"
        assert rows[0]["주소"] == "서울특별시 강남구 테헤란로 1"

    def test_gogaeksa_alias(self, tmp_path):
        """'고객사' 헤더도 회사명으로 인식."""
        p = tmp_path / "in.xlsx"
        _make_input(p, [
            ["고객사", "연락처", "주소"],
            ["현대모비스(주)", "041-599-9812", "경기도 평택시 포승읍"],
        ])
        rows = read_company_list(str(p))
        assert rows[0]["회사명"] == "현대모비스(주)"
        assert rows[0]["주소"] == "경기도 평택시 포승읍"
        assert rows[0]["전화번호"] == "041-599-9812"

    def test_fallback_keeps_other_matched_columns(self, tmp_path):
        """회사명 헤더만 인식 실패해도 주소 등 다른 컬럼은 보존."""
        p = tmp_path / "in.xlsx"
        _make_input(p, [
            ["알수없는헤더", "주소"],
            ["가나건설", "서울 강남구"],
        ])
        rows = read_company_list(str(p))
        assert rows[0]["회사명"] == "가나건설"
        assert rows[0]["주소"] == "서울 강남구"

    def test_no_header_falls_back_to_first_column(self, tmp_path):
        """헤더 인식 실패 시 1열을 회사명으로 처리."""
        p = tmp_path / "in.xlsx"
        _make_input(p, [
            ["abcdefg"],   # 매칭 안 됨 → 첫 데이터로 간주되지 않고 헤더 후보
            ["회사1"],
            ["회사2"],
        ])
        rows = read_company_list(str(p))
        # 헤더 매칭 실패 → 첫 컬럼을 회사명으로 보고 데이터 행만 채택 (헤더 행 1개는 건너뜀)
        assert all("회사명" in r for r in rows)
        # 첫 행이 헤더로 소비되었으므로 2개가 남음
        assert len(rows) == 2

    def test_skips_empty_company(self, tmp_path):
        p = tmp_path / "in.xlsx"
        _make_input(p, [
            ["회사명"],
            ["A"],
            [None],
            [""],
            ["B"],
        ])
        rows = read_company_list(str(p))
        assert [r["회사명"] for r in rows] == ["A", "B"]


class TestWriterInputColumns:
    def test_input_columns_left_and_show_once(self, tmp_path):
        """왼쪽에 [입력] 열, 동명 여러 건이면 입력값은 첫 행에만 표시."""
        p = tmp_path / "out.xlsx"
        inp = {"회사명": "동명건설", "주소": "서울"}
        write_results(
            str(p),
            records=[
                {"회사명": "동명건설(주)", "사업자번호": "111", "조회상태": "성공",
                 "_input": inp, "_input_show": True},
                {"회사명": "동명건설(주)", "사업자번호": "222", "조회상태": "성공",
                 "_input": inp, "_input_show": False},
            ],
            unfound=[], ambiguous=[], summary={},
        )
        wb = openpyxl.load_workbook(p)
        ws = wb["결과"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert headers[0] == "[입력] 회사명"
        assert headers[1] == "[입력] 주소"
        assert "회사명" in headers  # 결과 열도 존재
        # 첫 행엔 입력값, 둘째 행(동명 2번째)은 빈 칸
        assert ws.cell(2, 1).value == "동명건설"
        assert ws.cell(3, 1).value in (None, "")
        # 결과 열은 두 행 모두 채워짐
        name_col = headers.index("회사명") + 1
        assert ws.cell(2, name_col).value == "동명건설(주)"
        assert ws.cell(3, name_col).value == "동명건설(주)"


class TestResultSheetFiltering:
    def test_closed_and_unfound_rows_excluded_from_results(self, tmp_path):
        """결과 시트에는 성공(정상)만: 폐업자/휴업자·미발견·확인필요 제외.

        휴폐업 제외분은 미발견·오류 시트에 '제외' 사유로 남긴다.
        """
        p = tmp_path / "out.xlsx"
        write_results(
            str(p),
            records=[
                {"회사명": "정상사", "휴폐업정보": "일반과세자", "조회상태": "성공"},
                {"회사명": "폐업사", "휴폐업정보": "폐업자", "조회상태": "성공"},
                {"회사명": "휴업사", "휴폐업정보": "휴업자", "조회상태": "성공"},
                {"회사명": "못찾은사", "조회상태": "미발견"},
                {"회사명": "애매한사", "조회상태": "확인필요"},
            ],
            unfound=[{"회사명": "못찾은사", "조회상태": "미발견", "사유": "검색 0건"}],
            ambiguous=[], summary={},
        )
        wb = openpyxl.load_workbook(p)
        ws = wb["결과"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        name_col = headers.index("회사명") + 1
        names = [ws.cell(r, name_col).value for r in range(2, ws.max_row + 1)]
        assert names == ["정상사"]          # 결과 시트엔 정상만
        # 휴폐업 제외분은 미발견·오류 시트에 '제외'로 기록
        ws2 = wb["미발견·오류"]
        rows2 = [tuple(ws2.cell(r, c).value for c in (1, 2))
                 for r in range(2, ws2.max_row + 1)]
        assert ("폐업사", "제외") in rows2
        assert ("휴업사", "제외") in rows2
        assert ("못찾은사", "미발견") in rows2


class TestWriter:
    def test_creates_four_sheets(self, tmp_path):
        p = tmp_path / "out.xlsx"
        write_results(str(p), records=[], unfound=[], ambiguous=[], summary={})
        wb = openpyxl.load_workbook(p)
        assert wb.sheetnames == ["결과", "미발견·오류", "확인필요", "실행로그"]

    def test_writes_record_values(self, tmp_path):
        p = tmp_path / "out.xlsx"
        write_results(
            str(p),
            records=[{
                "회사명": "삼성전자", "대표자": "한종희", "사업자번호": "124-81-00998",
                "조회상태": "성공", "매출액(백만원)": 258900000,
            }],
            unfound=[{"회사명": "없는회사", "조회상태": "미발견", "사유": "검색 0건"}],
            ambiguous=[{"회사명": "동명", "채택후보": "A", "다른후보들": "B; C", "사유": "동점"}],
            summary={"total": 3, "success": 1, "not_found": 1, "ambiguous": 1, "error": 0,
                     "started_at": "2026-01-01 00:00:00"},
        )
        wb = openpyxl.load_workbook(p)
        ws = wb["결과"]
        # 헤더 + 1 데이터 행
        assert ws.cell(1, 1).value == "회사명"
        assert ws.cell(2, 1).value == "삼성전자"
        # 미발견 시트
        ws2 = wb["미발견·오류"]
        assert ws2.cell(2, 1).value == "없는회사"
        # 확인필요 시트
        ws3 = wb["확인필요"]
        assert ws3.cell(2, 2).value == "A"
        # 실행로그 시트에 통계 기록
        ws4 = wb["실행로그"]
        labels = [ws4.cell(r, 1).value for r in range(2, 9)]
        assert "총 건수" in labels
        assert "성공" in labels
