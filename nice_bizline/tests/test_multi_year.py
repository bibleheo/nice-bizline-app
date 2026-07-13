"""3개년 재무 옵션 동작 검증."""
import os
import openpyxl
import yaml

from nice_bizline.app.core.collector import MockCollector
from nice_bizline.app.core.worker import RunOptions, Worker, WorkerCallbacks
from nice_bizline.app.excelio.writer import build_headers, write_results


def _cfg():
    with open(os.path.join(os.path.dirname(__file__), "..", "config.yaml"),
              encoding="utf-8") as f:
        return yaml.safe_load(f)


def test_build_headers_single_year():
    h = build_headers([], finance_years=1)
    assert "매출액(백만원)" in h
    assert "영업이익(백만원)" in h
    assert "당기순이익(백만원)" in h


def test_build_headers_multi_year_collects_from_records():
    records = [
        {"매출액(2025)": 100, "매출액(2024)": 90, "영업이익(2025)": 10,
         "영업이익(2024)": 9, "당기순이익(2025)": 5, "당기순이익(2024)": 4},
        {"매출액(2023)": 80, "영업이익(2023)": 8, "당기순이익(2023)": 3},
    ]
    h = build_headers(records, finance_years=3)
    # 매출액 3개년이 모두 포함되고 최신부터 정렬
    sales_cols = [c for c in h if c.startswith("매출액")]
    assert sales_cols == ["매출액(2025)", "매출액(2024)", "매출액(2023)"]
    # 단년용 컬럼은 포함되지 않음
    assert "매출액(백만원)" not in h


def test_mock_collector_returns_finance_years_when_requested():
    collector = MockCollector(_cfg())
    candidates = collector.search("삼성전자")
    # 단년 모드
    detail1 = collector.fetch_detail(candidates[0], finance_years=1)
    assert "재무_연도별" not in detail1
    # 다년 모드
    detail3 = collector.fetch_detail(candidates[0], finance_years=3)
    assert "재무_연도별" in detail3
    assert len(detail3["재무_연도별"]) == 3
    assert detail3["재무_연도별"][0]["연도"] == "2025"


def test_worker_writes_multi_year_columns(tmp_path):
    cfg = _cfg()
    done = {}
    opts = RunOptions(
        user_id="u", password="p",
        companies=[{"회사명": "삼성전자"}, {"회사명": "현대자동차"}],
        finance_years=3,
    )
    cb = WorkerCallbacks(
        on_log=lambda *_: None,
        on_progress=lambda *_: None,
        on_done=done.update,
    )
    w = Worker(MockCollector(cfg), cfg, opts, cb)
    w.start()
    w.join(timeout=10)

    p = tmp_path / "out.xlsx"
    write_results(
        str(p),
        records=w.state.records,
        unfound=w.state.unfound,
        ambiguous=w.state.ambiguous,
        summary=done,
        finance_years=3,
    )

    wb = openpyxl.load_workbook(p)
    ws = wb["결과"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    # 3개년 매출액 컬럼이 모두 있어야 함
    for year in ("2025", "2024", "2023"):
        assert f"매출액({year})" in headers
        assert f"영업이익({year})" in headers
        assert f"당기순이익({year})" in headers
    # 단년 헤더는 없어야 함
    assert "매출액(백만원)" not in headers

    # 값 검증: 삼성전자 2025 매출액 (결과 회사명 컬럼은 [입력] 열들 뒤)
    name_col = headers.index("회사명") + 1
    samsung_row = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, name_col).value == "삼성전자(주)":
            samsung_row = r
            break
    assert samsung_row is not None
    sales_2025_col = headers.index("매출액(2025)") + 1
    assert ws.cell(samsung_row, sales_2025_col).value == 258900000
